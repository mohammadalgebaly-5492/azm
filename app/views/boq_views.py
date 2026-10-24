import io
import json
import secrets
from datetime import datetime, timedelta, timezone

from flask import Blueprint, abort, flash, g, redirect, render_template, request, send_file, url_for

from .. import boq as B
from ..auth import AuthError, login_required
from ..db import get_db, utcnow
from ..i18n import _, current_lang
from ..rbac import can, can_view_project

bp = Blueprint("boq", __name__)
# اقتراحات للوحدات حسب لغة الواجهة؛ الوحدة نص حر يُحفظ مع البند كما كُتب
UNITS = {
    "ar": ["م³", "م²", "م.ط", "طن", "كجم", "عدد", "مقطوعية", "لتر", "كيس", "يوم", "شهر"],
    "en": ["m³", "m²", "lm", "ton", "kg", "no.", "lump sum", "L", "bag", "day", "month"],
}


def _project(pid):
    db = get_db()
    p = db.execute("SELECT * FROM projects WHERE id = ?", (pid,)).fetchone()
    if p is None or not can_view_project(db, g.user, pid):
        abort(404)
    return db, p


def _manage():
    if not can(g.user, "boq.manage"):
        abort(403)


def _back(pid, anchor=""):
    return redirect(url_for("boq.page", pid=pid) + anchor)


@bp.get("/projects/<int:pid>/boq")
@login_required
def page(pid):
    db, p = _project(pid)
    nodes = B.wbs_nodes(db, pid)
    its = B.items(db, pid)
    by_wbs = {}
    for it in its:
        by_wbs.setdefault(it["wbs_id"], []).append(it)
    totals, total, by_type = B.summary(db, pid)
    checks = ok = None
    if not p["baseline_approved_at"]:
        checks, ok = B.baseline_checks(db, pid)[:2]
    approver = None
    if p["baseline_approved_by"]:
        approver = db.execute("SELECT full_name FROM users WHERE id = ?", (p["baseline_approved_by"],)).fetchone()
    return render_template("boq/page.html", p=p, nodes=nodes, by_wbs=by_wbs, totals=totals, total=total,
                           by_type=by_type, n_items=len(its), checks=checks, ok=ok, approver=approver,
                           locked=bool(p["baseline_approved_at"]), manage=can(g.user, "boq.manage"))


# ---------- WBS ----------

@bp.post("/projects/<int:pid>/wbs")
@login_required
def wbs_add(pid):
    db = _project(pid)[0]
    _manage()
    parent = request.form.get("parent_id")
    try:
        wid, code = B.add_wbs(db, pid, request.form.get("name"), int(parent) if parent and parent.isdigit() else None,
                              g.user["id"])
        flash(_("Section {code} added.", code=code), "ok")
    except AuthError as e:
        flash(str(e), "bad")
        return _back(pid)
    return _back(pid, f"#w{wid}")


@bp.post("/projects/<int:pid>/wbs/<int:wid>/rename")
@login_required
def wbs_rename(pid, wid):
    db = _project(pid)[0]
    _manage()
    try:
        B.rename_wbs(db, pid, wid, request.form.get("name"), g.user["id"])
        flash(_("Name updated."), "ok")
    except AuthError as e:
        flash(str(e), "bad")
    return _back(pid, f"#w{wid}")


@bp.post("/projects/<int:pid>/wbs/<int:wid>/delete")
@login_required
def wbs_delete(pid, wid):
    db = _project(pid)[0]
    _manage()
    try:
        B.delete_wbs(db, pid, wid, g.user["id"])
        flash(_("Section deleted."), "ok")
    except AuthError as e:
        flash(str(e), "bad")
    return _back(pid)


# ---------- البنود ----------

@bp.route("/projects/<int:pid>/boq/items/new", methods=["GET", "POST"])
@login_required
def item_new(pid):
    db, p = _project(pid)
    _manage()
    nodes = B.wbs_nodes(db, pid)
    if not nodes:
        flash(_("Add a WBS section first, then add items under it."), "bad")
        return _back(pid)
    error = None
    if request.method == "POST":
        wid = request.form.get("wbs_id", "")
        try:
            iid, code = B.add_item(db, pid, int(wid) if wid.isdigit() else 0, request.form, g.user["id"])
        except AuthError as e:
            error = str(e)
        else:
            flash(_("Item {code} added.", code=code), "ok")
            if request.form.get("again"):
                return redirect(url_for("boq.item_new", pid=pid, wbs=wid))
            return _back(pid, f"#i{iid}")
    return render_template("boq/item_form.html", p=p, it=None, nodes=nodes, form=request.form, error=error,
                           sel_wbs=request.values.get("wbs_id") or request.args.get("wbs"), UNITS=UNITS[current_lang()])


@bp.route("/projects/<int:pid>/boq/items/<int:iid>", methods=["GET", "POST"])
@login_required
def item_edit(pid, iid):
    db, p = _project(pid)
    _manage()
    it = db.execute("SELECT * FROM boq_items WHERE id = ? AND project_id = ?", (iid, pid)).fetchone()
    if it is None:
        abort(404)
    error = None
    if request.method == "POST":
        try:
            B.update_item(db, pid, iid, request.form, g.user["id"])
        except AuthError as e:
            error = str(e)
        else:
            flash(_("Item {code} saved.", code=it["code"]), "ok")
            return _back(pid, f"#i{iid}")
    form = request.form if error else {**dict(it), "qty": B.fmt_qty(it["qty_milli"]).replace(",", ""),
                                       "unit_rate": f"{it['unit_rate'] / 100:.2f}"}
    return render_template("boq/item_form.html", p=p, it=it, nodes=B.wbs_nodes(db, pid), form=form, error=error,
                           sel_wbs=str(it["wbs_id"]), UNITS=UNITS[current_lang()])


@bp.post("/projects/<int:pid>/boq/items/<int:iid>/delete")
@login_required
def item_delete(pid, iid):
    db = _project(pid)[0]
    _manage()
    try:
        B.delete_item(db, pid, iid, g.user["id"])
        flash(_("Item deleted."), "ok")
    except AuthError as e:
        flash(str(e), "bad")
    return _back(pid)


# ---------- الاستيراد والتصدير ----------

def _purge_previews(db):
    """حذف المعاينات الأقدم من ساعة."""
    cutoff = (datetime.now(timezone.utc) - timedelta(hours=1)).isoformat(timespec="seconds")
    db.execute("DELETE FROM import_previews WHERE created_at < ?", (cutoff,))


@bp.route("/projects/<int:pid>/boq/import", methods=["GET", "POST"])
@login_required
def import_upload(pid):
    db, p = _project(pid)
    _manage()
    nodes = B.wbs_nodes(db, pid)
    error = None
    if request.method == "POST":
        f = request.files.get("file")
        default_type = request.form.get("default_type") or "both"
        root = request.form.get("root_wbs", "")
        try:
            if not f or not f.filename:
                raise AuthError(_("Choose an Excel file."))
            rows = B.read_rows(f.filename, f.stream)
            parsed, errors = B.parse_import(rows, default_type)
        except AuthError as e:
            error = str(e)
        except Exception:
            error = _("The file couldn't be read. Make sure it's a valid Excel file (.xlsx), or start from the template.")
        else:
            # المعاينة تُحفظ في قاعدة البيانات لا على القرص: على Vercel القرص للقراءة فقط وكل طلب قد يصل لنسخة مختلفة
            token = secrets.token_urlsafe(16)
            _purge_previews(db)
            db.execute("INSERT INTO import_previews(token, project_id, user_id, root_wbs_id, payload, created_at) "
                       "VALUES (?,?,?,?,?,?)", (token, pid, g.user["id"], int(root) if root.isdigit() else None,
                                                json.dumps(parsed, ensure_ascii=False), utcnow()))
            total = sum(r["amount"] for r in parsed if r["kind"] == "item")
            return render_template("boq/import_preview.html", p=p, rows=parsed, errors=errors, token=token,
                                   total=total, filename=f.filename,
                                   root=next((n for n in nodes if str(n["id"]) == root), None))
    return render_template("boq/import.html", p=p, nodes=nodes, error=error, locked=bool(p["baseline_approved_at"]))


@bp.post("/projects/<int:pid>/boq/import/confirm")
@login_required
def import_confirm(pid):
    db = _project(pid)[0]
    _manage()
    token = request.form.get("token", "")
    if not token.replace("-", "").replace("_", "").isalnum():
        abort(400)
    _purge_previews(db)
    data = db.execute("DELETE FROM import_previews WHERE token = ? RETURNING project_id, user_id, root_wbs_id, payload",
                      (token,)).fetchone()
    if data is None:
        flash(_("The preview expired. Upload the file again."), "bad")
        return redirect(url_for("boq.import_upload", pid=pid))
    if data["project_id"] != pid or data["user_id"] != g.user["id"]:
        abort(403)
    try:
        n_items, n_heads = B.commit_import(db, pid, json.loads(data["payload"]), data["root_wbs_id"], g.user["id"])
    except AuthError as e:
        flash(str(e), "bad")
        return redirect(url_for("boq.import_upload", pid=pid))
    flash(_("Imported {items} items in {sections} sections.", items=n_items, sections=n_heads), "ok")
    return _back(pid)


@bp.get("/boq/template.xlsx")
@login_required
def template():
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOQ"
    ws.sheet_view.rightToLeft = current_lang() == "ar"
    head = [_("Item No"), _("Description"), _("Type"), _("Unit"), _("Qty"), _("Rate"), _("Amount")]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0D1CF2")
        c.alignment = Alignment(horizontal="center")
    both, supply = _(B.TYPES["both"][0]), _(B.TYPES["supply"][0])
    m3, m2, ton = (UNITS[current_lang()][i] for i in (0, 1, 3))
    sample = [
        ["1", _("Concrete works"), None, None, None, None, None],
        ["1.1", _("Plain concrete under footings C15, 10 cm thick"), both, m3, 120, 280, None],
        ["1.2", _("Reinforced concrete for footings C35 incl. rebar and formwork"), both, m3, 850, 1150, None],
        ["2", _("Masonry works"), None, None, None, None, None],
        ["2.1", _("Hollow concrete block 20 cm for external walls"), both, m2, 2400, 95, None],
        ["3", _("Supplies"), None, None, None, None, None],
        ["3.1", _("Supply of rebar, 16 mm diameter"), supply, ton, 180, 2750, None],
    ]
    for r in sample:
        ws.append(r)
        n = ws.max_row
        if r[4] is not None:
            ws.cell(n, 7, f"=E{n}*F{n}")
        else:
            ws.cell(n, 2).font = Font(bold=True)
    for col, w in zip("ABCDEFG", (11, 60, 16, 10, 12, 14, 16)):
        ws.column_dimensions[col].width = w
    notes = wb.create_sheet(_("Instructions"))
    notes.sheet_view.rightToLeft = current_lang() == "ar"
    for line in (_("The first row holds the column headers: Item No, Description, Type, Unit, Qty, Rate, Amount."),
                 _("A row with a description but no quantity or rate is a section heading. It becomes a WBS section, "
                   "and the items after it are filed under it."),
                 _("The Type column is optional: Supply, Execution, Supply & Execution. If empty, the default type "
                   "you choose during import is used."),
                 _("The Amount column is optional. If present, the platform checks it equals quantity × rate."),
                 _("Arabic and English column headers are both accepted.")):
        notes.append([line])
    notes.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="azm-boq-template.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.get("/projects/<int:pid>/boq/export.xlsx")
@login_required
def export(pid):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    db, p = _project(pid)
    nodes = B.wbs_nodes(db, pid)
    by_wbs = {}
    for it in B.items(db, pid):
        by_wbs.setdefault(it["wbs_id"], []).append(it)
    totals, total = B.summary(db, pid)[:2]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOQ"
    ws.sheet_view.rightToLeft = current_lang() == "ar"
    ws.append([f"{p['code']} — {p['name']}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([_("Code"), _("Item No"), _("Description"), _("Type"), _("Unit"), _("Qty"), _("Rate"), _("Amount")])
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0D1CF2")
    for n in nodes:
        ws.append([n["code"], None, n["name"], None, None, None, None, totals[n["id"]][0] / 100])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
        for it in by_wbs.get(n["id"], []):
            ws.append([it["code"], it["ref_no"], it["description"], _(B.TYPES[it["item_type"]][0]), it["unit"],
                       it["qty_milli"] / 1000, it["unit_rate"] / 100, it["amount"] / 100])
    ws.append([None, None, _("Grand total"), None, None, None, None, total / 100])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=3, min_col=6, max_col=8):
        for c in row:
            c.number_format = "#,##0.00"
    for col, w in zip("ABCDEFGH", (18, 10, 60, 14, 10, 12, 14, 18)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{p['code']}-BOQ.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------- خط الأساس ----------

@bp.post("/projects/<int:pid>/baseline/approve")
@login_required
def baseline_approve(pid):
    db = _project(pid)[0]
    _manage()
    if request.form.get("confirm") != "1":
        flash(_("Tick the confirmation box before approving."), "bad")
        return _back(pid, "#baseline")
    try:
        total = B.approve_baseline(db, pid, g.user["id"])
        flash(_("Baseline approved at {amount} SAR. The BOQ is now frozen.", amount=f"{total / 100:,.2f}"), "ok")
    except AuthError as e:
        flash(str(e), "bad")
    return _back(pid, "#baseline")
