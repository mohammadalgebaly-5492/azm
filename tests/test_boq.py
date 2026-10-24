"""اختبارات المرحلة الثانية: WBS وجدول الكميات والاستيراد وخط الأساس."""
import io
import os
import re
import sqlite3
import tempfile
import unittest

import openpyxl

from app import boq as B
from app import create_app
from app.auth import AuthError, create_user
from app.db import connect
from app.projects import add_member, create_project


class BoqBase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.app = create_app({"DATABASE": os.path.join(self.tmp.name, "b.sqlite3"), "SESSION_COOKIE_SECURE": False,
                               "SECRET_KEY": "t"})
        self.db = connect(self.app.config["DATABASE"])
        with self.app.app_context():
            self.admin, _, _ = create_user(self.db, username="admin", email="a@co.sa", full_name="الأدمن",
                                           role="admin", password="AdminPass123")
            self.pmid, _, _ = create_user(self.db, username="pm.one", email="pm@co.sa", full_name="مدير", role="pm",
                                          password="PmPass12345")
            self.pid, _ = create_project(self.db, {"name": "مشروع مسك", "contract_value": "1000000",
                                                   "start_date": "2026-11-01"}, self.admin)
            add_member(self.db, self.pid, self.pmid, self.admin)
        self.c = self.app.test_client()
        self.c.post("/login", data={"username": "admin", "password": "AdminPass123"})

    def tearDown(self):
        self.db.close()
        self.tmp.cleanup()

    def csrf(self, cl=None):
        html = (cl or self.c).get("/").get_data(as_text=True)
        return re.search(r'name="csrf_token" value="([^"]+)"', html).group(1)


class Codes(BoqBase):
    def test_wbs_and_item_codes(self):
        with self.app.app_context():
            a, ca = B.add_wbs(self.db, self.pid, "أعمال الخرسانة", None, self.admin)
            b, cb = B.add_wbs(self.db, self.pid, "أعمال المباني", None, self.admin)
            a1, ca1 = B.add_wbs(self.db, self.pid, "القواعد", a, self.admin)
            self.assertEqual((ca, cb, ca1), ("01", "02", "01.01"))
            _, i1 = B.add_item(self.db, self.pid, a1, {"description": "خرسانة مسلحة C35", "unit": "م³",
                                                        "qty": "850.5", "unit_rate": "1150", "item_type": "both"}, self.admin)
            _, i2 = B.add_item(self.db, self.pid, a1, {"description": "توريد حديد Ø16", "unit": "طن",
                                                        "qty": "180", "unit_rate": "2750", "item_type": "supply"}, self.admin)
            self.assertEqual((i1, i2), ("01.01-SE-010", "01.01-S-020"))

    def test_amount_rounding_and_db_check(self):
        self.assertEqual(B.amount_of(B.parse_qty("850.5"), B.parse_rate("1150")), 97807500)
        self.assertEqual(B.amount_of(B.parse_qty("0.333"), B.parse_rate("1.00")), 33)
        self.assertEqual(B.parse_qty("١٬٢٥٠٫٥"), 1250500)
        with self.assertRaises(AuthError):
            B.parse_qty("0")
        with self.app.app_context():
            w, _ = B.add_wbs(self.db, self.pid, "قسم", None, self.admin)
            with self.assertRaises(sqlite3.IntegrityError):  # الإجمالي لازم يساوي الكمية × السعر
                self.db.execute("INSERT INTO boq_items(project_id, wbs_id, code, description, item_type, unit, qty_milli,"
                                " unit_rate, amount, created_at, created_by) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                                (self.pid, w, "X", "تلاعب", "both", "م", 1000, 100, 999, "t", self.admin))

    def test_wbs_depth_limit_and_delete_rules(self):
        with self.app.app_context():
            p = None
            for i in range(4):
                p, _ = B.add_wbs(self.db, self.pid, f"مستوى {i + 1}", p, self.admin)
            with self.assertRaises(AuthError):
                B.add_wbs(self.db, self.pid, "مستوى 5", p, self.admin)
            B.add_item(self.db, self.pid, p, {"description": "بند تجريبي", "unit": "م²", "qty": "1",
                                              "unit_rate": "1", "item_type": "exec"}, self.admin)
            with self.assertRaises(AuthError):
                B.delete_wbs(self.db, self.pid, p, self.admin)


class ImportTests(BoqBase):
    def _xlsx(self, rows):
        wb = openpyxl.Workbook()
        for r in rows:
            wb.active.append(r)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def test_headings_become_wbs_and_mismatch_flagged(self):
        rows = [["جدول كميات مشروع مسك"], [],
                ["رقم البند", "الوصف", "الوحدة", "الكمية", "سعر الوحدة", "الإجمالي"],
                ["1", "أعمال الخرسانة", None, None, None, None],
                ["1.1", "خرسانة عادية C15", "م³", 120, 280, 33600],
                ["1.2", "خرسانة مسلحة C35", "م³", 850, 1150, 977500],
                ["2", "أعمال المباني", None, None, None, None],
                ["2.1", "بلوك 20 سم", "م²", 2400, 95.5, 1],
                ["2.2", "بلوك 15 سم", "م²", 600, "٨٥", None]]
        parsed, errors = B.parse_import(rows, "both")
        self.assertEqual([r["kind"] for r in parsed], ["heading", "item", "item", "heading", "item"])
        self.assertEqual(len(errors), 1)
        self.assertEqual(errors[0][0], 8)

        r = self.c.post(f"/projects/{self.pid}/boq/import", data={
            "csrf_token": self.csrf(), "default_type": "both", "root_wbs": "",
            "file": (self._xlsx(rows), "boq.xlsx")}, content_type="multipart/form-data")
        html = r.get_data(as_text=True)
        self.assertIn("معاينة الاستيراد", html)
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM boq_items").fetchone()[0], 0)  # لا حفظ قبل التأكيد
        token = re.search(r'name="token" value="([^"]+)"', html).group(1)
        self.c.post(f"/projects/{self.pid}/boq/import/confirm", data={"csrf_token": self.csrf(), "token": token})
        codes = [r[0] for r in self.db.execute("SELECT code FROM wbs_nodes ORDER BY code")]
        self.assertEqual(codes, ["01", "02"])
        self.assertEqual(tuple(self.db.execute("SELECT COUNT(*), SUM(amount) FROM boq_items").fetchone()),
                         (3, (33600 + 977500 + 51000) * 100))
        # نفس المعاينة لا تُستخدم مرتين
        self.c.post(f"/projects/{self.pid}/boq/import/confirm", data={"csrf_token": self.csrf(), "token": token})
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM boq_items").fetchone()[0], 3)

    def test_english_headers_and_types(self):
        rows = [["Item No", "Description", "Type", "Unit", "Qty", "Rate"],
                ["A", "Rebar supply", "Supply", "ton", 10, 2750],
                ["B", "Waterproofing", "تنفيذ", "m2", 400, 55]]
        parsed, errors = B.parse_import(rows, "both")
        self.assertFalse(errors)
        self.assertEqual([r["item_type"] for r in parsed], ["supply", "exec"])

    def test_missing_headers_rejected(self):
        with self.assertRaises(AuthError):
            B.parse_import([["a", "b"], [1, 2]], "both")


class BaselineAndPermissions(BoqBase):
    def _seed(self):
        with self.app.app_context():
            w, _ = B.add_wbs(self.db, self.pid, "أعمال الخرسانة", None, self.admin)
            iid, _ = B.add_item(self.db, self.pid, w, {"description": "خرسانة", "unit": "م³", "qty": "100",
                                                        "unit_rate": "1000", "item_type": "both"}, self.admin)
        return w, iid

    def test_baseline_freezes_boq(self):
        w, iid = self._seed()
        r = self.c.post(f"/projects/{self.pid}/baseline/approve", data={"csrf_token": self.csrf(), "confirm": "1"})
        self.assertEqual(r.status_code, 302)
        p = self.db.execute("SELECT baseline_total, status FROM projects WHERE id = ?", (self.pid,)).fetchone()
        self.assertEqual(tuple(p), (10000000, "active"))
        with self.app.app_context():
            for fn in (lambda: B.add_item(self.db, self.pid, w, {"description": "جديد", "unit": "م", "qty": "1",
                                                                 "unit_rate": "1", "item_type": "exec"}, self.admin),
                       lambda: B.delete_item(self.db, self.pid, iid, self.admin),
                       lambda: B.add_wbs(self.db, self.pid, "قسم جديد", None, self.admin),
                       lambda: B.approve_baseline(self.db, self.pid, self.admin)):
                with self.assertRaises(AuthError):
                    fn()
        with self.assertRaises(sqlite3.IntegrityError):  # التجميد مفروض في قاعدة البيانات نفسها
            self.db.execute("UPDATE boq_items SET unit_rate = 1, amount = 100 WHERE id = ?", (iid,))

    def test_baseline_blocked_without_items(self):
        r = self.c.post(f"/projects/{self.pid}/baseline/approve", data={"csrf_token": self.csrf(), "confirm": "1"},
                        follow_redirects=True)
        self.assertIn("لا يمكن الاعتماد", r.get_data(as_text=True))

    def test_pm_reads_but_cannot_edit(self):
        w, iid = self._seed()
        pm = self.app.test_client()
        self.db.execute("UPDATE users SET must_change_password = 0")
        pm.post("/login", data={"username": "pm.one", "password": "PmPass12345"})
        page = pm.get(f"/projects/{self.pid}/boq")
        self.assertEqual(page.status_code, 200)
        self.assertNotIn("بند جديد", page.get_data(as_text=True))
        tok = self.csrf(pm)
        self.assertEqual(pm.get(f"/projects/{self.pid}/boq/items/new").status_code, 403)
        self.assertEqual(pm.post(f"/projects/{self.pid}/boq/items/{iid}/delete", data={"csrf_token": tok}).status_code, 403)
        self.assertEqual(pm.post(f"/projects/{self.pid}/baseline/approve",
                                 data={"csrf_token": tok, "confirm": "1"}).status_code, 403)
        self.assertEqual(pm.post(f"/projects/{self.pid}/wbs", data={"csrf_token": tok, "name": "دخيل"}).status_code, 403)
        self.assertEqual(pm.get(f"/projects/{self.pid}/boq/export.xlsx").status_code, 200)

    def test_non_member_gets_404(self):
        with self.app.app_context():
            create_user(self.db, username="other", email="o@co.sa", full_name="غريب", role="pm", password="OtherPass123")
        self.db.execute("UPDATE users SET must_change_password = 0")
        cl = self.app.test_client()
        cl.post("/login", data={"username": "other", "password": "OtherPass123"})
        self.assertEqual(cl.get(f"/projects/{self.pid}/boq").status_code, 404)
        self.assertEqual(cl.get(f"/projects/{self.pid}/boq/export.xlsx").status_code, 404)

    def test_item_form_and_template_download(self):
        w, _ = self._seed()
        r = self.c.post(f"/projects/{self.pid}/boq/items/new", data={
            "csrf_token": self.csrf(), "wbs_id": w, "description": "بلوك 20 سم", "unit": "م²", "qty": "2,400",
            "unit_rate": "95", "item_type": "both"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.db.execute("SELECT amount FROM boq_items WHERE description = 'بلوك 20 سم'").fetchone()[0],
                         22800000)
        t = self.c.get("/boq/template.xlsx")
        self.assertEqual(t.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(t.data))
        parsed, errors = B.parse_import([list(r) for r in wb.active.iter_rows(values_only=True)], "both")
        self.assertFalse(errors)
        self.assertEqual(sum(1 for x in parsed if x["kind"] == "item"), 4)


if __name__ == "__main__":
    unittest.main()
